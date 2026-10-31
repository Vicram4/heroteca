"""
Vuelca el contenido de los .docx/.pptx a texto legible para analizar su estructura.

Uso:
    .venv/Scripts/python.exe scripts/inspeccionar.py <carpeta_originales> <carpeta_salida>

Los originales solo se leen. Por cada documento se crea un .txt en la salida
(con la misma ruta relativa) y un resumen general en `inventario.tsv`.
"""
import sys
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.util import Emu

NS_IMG = '{http://schemas.openxmlformats.org/drawingml/2006/main}blip'


def bloques_docx(doc):
    """Recorre párrafos y tablas en el orden en que aparecen."""
    for hijo in doc.element.body.iterchildren():
        if hijo.tag.endswith('}p'):
            yield Paragraph(hijo, doc)
        elif hijo.tag.endswith('}tbl'):
            yield Table(hijo, doc)


def texto_celda(celda):
    return ' / '.join(p.text.strip() for p in celda.paragraphs if p.text.strip())


def volcar_docx(ruta):
    doc = Document(ruta)
    lineas, n_par, n_tab = [], 0, 0
    n_img = len(list(doc.element.body.iter(NS_IMG)))
    for b in bloques_docx(doc):
        if isinstance(b, Paragraph):
            t = b.text.strip()
            if not t:
                continue
            n_par += 1
            estilo = b.style.name if b.style is not None else '?'
            negrita = any(r.bold for r in b.runs if r.text.strip())
            lineas.append(f'[{estilo}{" ·N" if negrita else ""}] {t}')
        else:
            n_tab += 1
            filas = b.rows
            lineas.append(f'--- TABLA {n_tab}: {len(filas)} filas x {len(b.columns)} columnas')
            for fila in filas:
                celdas = []
                for c in fila.cells:
                    txt = texto_celda(c)
                    if not celdas or celdas[-1] != txt:  # celdas combinadas se repiten
                        celdas.append(txt)
                lineas.append('  | ' + ' | '.join(celdas))
            lineas.append('--- FIN TABLA')
    sec = doc.sections[0]
    cab = f'Páginas: {"horizontal" if sec.page_width > sec.page_height else "vertical"}; párrafos: {n_par}; tablas: {n_tab}; imágenes: {n_img}'
    return cab, lineas, (n_par, n_tab, n_img)


def volcar_pptx(ruta):
    prs = Presentation(ruta)
    lineas, n_img, n_tab = [], 0, 0

    def recorrer(formas, sangria='  '):
        nonlocal n_img, n_tab
        for f in formas:
            pos = ''
            if f.left is not None and f.top is not None:
                pos = f' @({Emu(f.left).cm:.0f},{Emu(f.top).cm:.0f})cm'
            if f.shape_type == 6:  # grupo
                lineas.append(f'{sangria}<grupo{pos}>')
                recorrer(f.shapes, sangria + '  ')
            elif f.shape_type == 13:  # imagen
                n_img += 1
                lineas.append(f'{sangria}<imagen{pos}> {f.name}')
            elif getattr(f, 'has_table', False) and f.has_table:
                n_tab += 1
                lineas.append(f'{sangria}<tabla{pos}>')
                for fila in f.table.rows:
                    lineas.append(sangria + '  | ' + ' | '.join(c.text.strip().replace('\n', ' / ') for c in fila.cells))
            elif f.has_text_frame and f.text_frame.text.strip():
                txt = f.text_frame.text.strip().replace('\n', ' / ').replace('\x0b', ' / ')
                lineas.append(f'{sangria}<texto{pos}> {txt}')

    for i, s in enumerate(prs.slides, 1):
        lineas.append(f'=== DIAPOSITIVA {i} (diseño: {s.slide_layout.name})')
        recorrer(s.shapes)
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
            lineas.append('  <notas> ' + s.notes_slide.notes_text_frame.text.strip().replace('\n', ' / '))
    ancho, alto = Emu(prs.slide_width).cm, Emu(prs.slide_height).cm
    cab = f'Diapositivas: {len(prs.slides)}; tamaño: {ancho:.1f}x{alto:.1f} cm; imágenes: {n_img}; tablas: {n_tab}'
    return cab, lineas, (len(prs.slides), n_tab, n_img)


def main(origen, salida):
    origen, salida = Path(origen), Path(salida)
    inventario = ['ruta\ttipo\tcantidad\ttablas\timagenes']
    for ruta in sorted(origen.rglob('*')):
        ext = ruta.suffix.lower()
        if ext not in ('.docx', '.pptx') or ruta.name.startswith('~$'):
            continue
        rel = ruta.relative_to(origen)
        try:
            cab, lineas, (n, t, im) = (volcar_docx if ext == '.docx' else volcar_pptx)(ruta)
        except Exception as e:  # noqa: BLE001 — se anota y se sigue
            cab, lineas, (n, t, im) = f'ERROR: {e}', [], (0, 0, 0)
        destino = salida / rel.with_suffix(rel.suffix + '.txt')
        destino.parent.mkdir(parents=True, exist_ok=True)
        destino.write_text(f'{rel}\n{cab}\n\n' + '\n'.join(lineas) + '\n', encoding='utf-8')
        inventario.append(f'{rel}\t{ext[1:]}\t{n}\t{t}\t{im}')
    (salida / 'inventario.tsv').write_text('\n'.join(inventario) + '\n', encoding='utf-8')
    print(f'{len(inventario) - 1} documentos volcados en {salida}')


if __name__ == '__main__':
    main(*sys.argv[1:3])
