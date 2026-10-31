"""
Catálogo de DC a partir de las copias de documentos/originales/dc (nunca las modifica).

DC se organiza como el disco duro del usuario:
  - TODO UDC/UDC Películas por orden/001 UDC   -> películas clásicas (n.º = PRE-DCUE, orden de estreno)
  - TODO UDC/UDC Películas por orden/002 DCDU  -> universo extendido (n.º = orden de la historia)
  - UDC Series por orden                        -> series (n.º de carpeta = estreno; códigos 1.1, 2.3… = visionado)
  - animacion/Animación/001-036                 -> películas animadas
Los Word del Archivador AZUL completan año, calidad, carátulas y capítulos.

Uso:  python scripts/dc_catalogo.py documentos/originales/dc documentos/extraido/dc
Salida: catalogo.json, imagenes/ (carátulas de los índices) y avisos.txt
"""
import hashlib
import json
import os
import re
import sys
from pathlib import Path

import docx
from docx.oxml.ns import qn

INSIGNIAS = {
    '971dc0bfe2': '4k',
    'fb77ad2f9b': '1080p',
    '9ba855e36f': '1080p',
    '0b38fcd8a9': '720p',
    '0940f1b103': 'avi',
    'f59225962c': 'webm',
}
VACIA = '189fab5fa7'  # recuadro punteado: sin copia
LOGO_DC = 'ff2233a4b0'

# Abreviaturas de las series en el orden capítulo a capítulo de los Crossover
ABREV = {
    'tflash': 'The Flash', 'flash': 'The Flash', 'arrow': 'Arrow', 'sprgirl': 'Supergirl', 'supergirl': 'Supergirl',
    'legen': "Legends of Tomorrow", 'legends': "Legends of Tomorrow", 'vixen': 'Vixen', 'cisco': 'Crónicas de Cisco',
    'ray': 'The Ray', 'theray': 'The Ray', 'blight': 'Black Lightning', 'black': 'Black Lightning', 'batw': 'Batwoman',
    'batwoman': 'Batwoman', 'consta': 'Constantine',
}

avisos = []


def limpiar(t):
    return re.sub(r'\s+', ' ', (t or '').replace(' ', ' ')).strip()


def textos(elemento):
    """Párrafos no vacíos (incluidos los de los cuadros de texto) en orden de documento."""
    salida = []
    for p in elemento.iter(qn('w:p')):
        # Un párrafo de un cuadro de texto aparece dentro de otro párrafo: solo el texto propio
        propio = ''.join(t.text or '' for t in p.iter(qn('w:t')) if _padre_p(t) is p)
        if limpiar(propio):
            salida.append(limpiar(propio))
    return salida


def _padre_p(nodo):
    n = nodo.getparent()
    while n is not None and n.tag != qn('w:p'):
        n = n.getparent()
    return n


def imagenes(elemento, rels):
    salida = []
    for b in elemento.iter(qn('a:blip')):
        rid = b.get(qn('r:embed'))
        if rid in rels:
            blob = rels[rid].target_part.blob
            salida.append({'hash': hashlib.sha1(blob).hexdigest()[:10], 'blob': blob,
                           'ext': os.path.splitext(rels[rid].target_ref)[1].lower() or '.png'})
    return salida


def guardar_imagen(img, carpeta, nombre):
    carpeta.mkdir(parents=True, exist_ok=True)
    ruta = carpeta / f'{nombre}{img["ext"]}'
    ruta.write_bytes(img['blob'])
    return ruta.as_posix()


ETIQUETAS = re.compile(r'^\((.+)\)$')


def leer_pagina_indice(ruta, salida_img, prefijo):
    """Páginas «Películas por orden de estreno» (PRE-DCUE / DCUE) y «pelis solo»."""
    d = docx.Document(ruta)
    rels = d.part.rels
    cabecera = ' / '.join(textos(d.tables[0]._tbl)) if d.tables else ''
    filas = []
    for fila in d.tables[-1].rows:
        celdas = []
        for c in fila.cells:  # las celdas combinadas se repiten: quitarlas
            if not celdas or c._tc is not celdas[-1]._tc:
                celdas.append(c)
        txt = []
        for c in celdas:
            for t in textos(c._tc):
                if t not in txt:
                    txt.append(t)
        if not any(re.match(r'^\d{3}', t) for t in txt):
            continue
        imgs = [imagenes(c._tc, rels) for c in celdas]
        e = {'pagina': Path(ruta).name, 'etiquetas': [], 'calidad': []}
        resto = []
        for t in txt:
            m = re.match(r'^(\d{3})\s*(?:UDC)?\s*(?:\((\d{4})\))?\s*$', t)
            if m and 'numero' not in e:
                e['numero'] = int(m.group(1))
                if m.group(2):
                    e['anio'] = int(m.group(2))
            elif re.fullmatch(r'\(?\d{4}\)?', t) and 'anio' not in e:
                e['anio'] = int(t.strip('()'))
            elif ETIQUETAS.match(t):
                e['etiquetas'].append(ETIQUETAS.match(t).group(1))
            elif t in ('Ficha técnica', '.', 'B.S.O.'):
                if t == 'B.S.O.':
                    e['etiquetas'].append('B.S.O.')
            else:
                resto.append(t)
        if resto:
            e['titulo'] = resto[0].rstrip(':').strip()
            if len(resto) > 1:
                e['subtitulo'] = ' '.join(resto[1:]).rstrip('.').strip()
                e['titulo'] = resto[0].strip()
        planas = [i for grupo in imgs for i in grupo]
        for i in planas:
            if i['hash'] in INSIGNIAS and INSIGNIAS[i['hash']] not in e['calidad']:
                e['calidad'].append(INSIGNIAS[i['hash']])
        if imgs and imgs[0]:
            e['portada'] = guardar_imagen(imgs[0][0], salida_img, f'{prefijo}-{e.get("numero", 0):03d}')
        if len(imgs) > 1 and imgs[-1] and imgs[-1][0]['hash'] not in INSIGNIAS:
            e['logo'] = guardar_imagen(imgs[-1][0], salida_img, f'{prefijo}-{e.get("numero", 0):03d}-logo')
        filas.append(e)
    return cabecera, filas


# --- Series -----------------------------------------------------------------

CAP = re.compile(r'^(\d{1,2}x\d{1,3})\s*-+\s*(.+?)[.-]*$')
CAP_CROSS = re.compile(r'^(\d{1,3})\s*\+?\s*([A-Za-z ]+?)\s*(\d{1,2}x\d{1,3})\s*-?\s*(.+?)\.?$')
BLOQUE = re.compile(r'^(\d{1,2}\.\d)\s+(.+)$')


def nombre_serie(abrev):
    k = re.sub(r'[^a-z]', '', abrev.lower())
    completos = {re.sub(r'[^a-z]', '', v.lower()): v for v in ABREV.values()}
    if k in completos:
        return completos[k]
    if k not in ABREV:
        avisos.append(f'Abreviatura de serie desconocida en Crossover: «{abrev}»')
    return ABREV.get(k, abrev.strip())


def leer_pagina_series(ruta, salida_img):
    d = docx.Document(ruta)
    rels = d.part.rels
    bloques = []
    for fila in d.tables[-1].rows:
        celdas = []
        for c in fila.cells:
            if not celdas or c._tc is not celdas[-1]._tc:
                celdas.append(c)
        cols = [textos(c._tc) for c in celdas]
        if not cols or not cols[0]:
            continue
        m = BLOQUE.match(cols[0][0]) or re.match(r'^(\d{1,2}\.\d)()$', cols[0][0])
        if not m:
            continue
        codigo, nombre = m.group(1), limpiar(m.group(2)) or (cols[0][1] if len(cols[0]) > 1 else '')
        intercalado = any(CAP_CROSS.match(t) for col in cols[1:-1] for t in col)
        if intercalado or re.match(r'^(MEGA)?CROSS[OA]WER', nombre.upper()):
            if not re.match(r'^(MEGA)?CROSS', nombre.upper()):  # 4.3: «Arrow / Temporada 3 completa / The Flash…»
                cols[0] = [cols[0][0]] + cols[0]
            series = []
            for col in (cols[0], cols[-1]):
                for i, t in enumerate(col[1:], 1):
                    mt = re.match(r'^Temporada\s+(\d+)', t)
                    if mt and i > 0 and not re.fullmatch(r'\d', col[i - 1]):
                        series.append({'serie': re.sub(r'^\d{1,2}\.\d\s+', '', col[i - 1]), 'temporada': int(mt.group(1))})
            portadas = [i for c in (celdas[0], celdas[-1]) for i in imagenes(c._tc, rels) if i['hash'] != LOGO_DC]
            for k, (sr, im) in enumerate(zip(series, portadas)):
                sr['portada'] = guardar_imagen(im, salida_img, f"serie-{codigo}-{Path(ruta).stem[-3:]}-{k + 1}")
            orden = []
            pendiente = None  # «29Megacrossower Supergirl» + en la línea siguiente «2x08-Medusa.»
            for col in cols[1:-1]:
                for t in col:
                    mc = CAP_CROSS.match(t)
                    mm = re.match(r'^(\d{1,3})\s*Megacross[oa]wer\s+(.+)$', t, re.I)
                    ma = re.match(r'^(\d{1,3})\s+(.+?)\s*\(animación\)$', t, re.I)
                    mb = re.match(r'^(.+?)\s+(\d{1,2}x\d{2,3})\s*-\s*"?(.+?)"?:?$', t)  # borrador sin numerar
                    if pendiente and CAP.match(t):
                        mcap = CAP.match(t)
                        orden.append({**pendiente, 'codigo': mcap.group(1), 'titulo': limpiar(mcap.group(2))})
                        pendiente = None
                    elif mc:
                        orden.append({'n': int(mc.group(1)), 'serie': nombre_serie(mc.group(2)),
                                      'codigo': mc.group(3), 'titulo': limpiar(mc.group(4))})
                    elif mm:
                        pendiente = {'n': int(mm.group(1)), 'serie': nombre_serie(mm.group(2)), 'megacrossover': True}
                    elif re.match(r'^(\d{1,3})\s+(Crónicas de CISCO|[A-Z][\w ]+?)\s*$', t, re.I) and not re.search(r'\dx\d', t):
                        mw = re.match(r'^(\d{1,3})\s+(.+?)\s*$', t)
                        orden.append({'n': int(mw.group(1)), 'serie': nombre_serie(mw.group(2)), 'codigo': None,
                                      'titulo': None, 'completa': True})
                    elif ma:
                        orden.append({'n': int(ma.group(1)), 'serie': nombre_serie(ma.group(2)),
                                      'codigo': None, 'titulo': None, 'animada': True})
                    elif mb and re.sub(r'[^a-z]', '', mb.group(1).lower()) in {re.sub(r'[^a-z]', '', k.lower()) for k in ABREV.values()} | set(ABREV):
                        orden.append({'n': None, 'serie': nombre_serie(mb.group(1)), 'codigo': mb.group(2),
                                      'titulo': limpiar(mb.group(3)), 'borrador': True})
                    elif not re.fullmatch(r'[-\s]*Temporada.*|\(.*\)|.*Crónicas.*|.*completa.*', t, re.I):
                        avisos.append(f'{Path(ruta).name} {codigo}: línea no reconocida en Crossover: «{t}»')
            bloques.append({'codigo': codigo, 'tipo': 'crossover', 'series': series, 'orden': orden,
                            'pagina': Path(ruta).name})
        else:
            temp = next((int(x) for t in cols[0][1:] for x in re.findall(r'Temporada\s+(\d+)', t)), None)
            caps = []
            for col in cols[1:]:
                for t in col:
                    mc = CAP.match(t)
                    if mc:
                        caps.append({'codigo': mc.group(1), 'titulo': limpiar(mc.group(2))})
            b = {'codigo': codigo, 'tipo': 'temporada', 'serie': nombre, 'temporada': temp,
                 'capitulos': caps, 'pagina': Path(ruta).name}
            im = [i for i in imagenes(celdas[0]._tc, rels) if i['hash'] != LOGO_DC]
            if im:
                b['portada'] = guardar_imagen(im[0], salida_img, f'serie-{codigo}-{Path(ruta).stem[-3:]}')
            bloques.append(b)
    return bloques


def leer_resumen(ruta):
    d = docx.Document(ruta)
    carpetas, espera = [], []
    for fila in d.tables[-1].rows:
        vistos = []
        for c in fila.cells:
            if vistos and c._tc is vistos[-1]._tc:
                continue
            vistos.append(c)
            t = textos(c._tc)
            if not t:
                continue
            m = re.match(r'^Carpeta\s+(\d+)', t[0])
            if m:
                carpetas.append({'carpeta': int(m.group(1)), 'etapa': t[1] if len(t) > 1 else '', 'contenido': t[2:]})
            elif t[0].upper().startswith('EN ESPERA'):
                espera += t[1:]
            elif carpetas == [] or espera:
                espera += t
    return carpetas, espera


# --- Disco ------------------------------------------------------------------

def leer_estructura(ruta):
    """Lee ESTRUCTURA-DEL-DISCO.txt (árbol con sangría de 2 espacios)."""
    pila, nodos = [], []
    for linea in Path(ruta).read_text(encoding='utf-8').splitlines()[1:]:
        nivel = (len(linea) - len(linea.lstrip(' '))) // 2
        nombre = linea.strip()
        pila = pila[:nivel] + [nombre.rstrip('/')]
        nodos.append({'ruta': '/'.join(pila), 'nivel': nivel, 'carpeta': nombre.endswith('/'), 'nombre': nombre.rstrip('/')})
    return nodos


def hijos(nodos, ruta, solo_carpetas=True):
    n = ruta.count('/') + 1
    return [x for x in nodos if x['ruta'].startswith(ruta + '/') and x['nivel'] == n and (x['carpeta'] or not solo_carpetas)]


NUM_TITULO_ANIO = re.compile(r'^(\d{3})\s+(.+?)(?:\s+(\d{4}))?(?:\s+\((.+)\))?$')


def peliculas_disco(nodos, ruta):
    salida = []
    for x in hijos(nodos, ruta):
        m = NUM_TITULO_ANIO.match(x['nombre'])
        if m:
            salida.append({'numero': int(m.group(1)), 'carpeta': x['nombre'], 'titulo': m.group(2).strip(),
                           'anio': int(m.group(3)) if m.group(3) else None, 'etiqueta': m.group(4)})
        else:
            salida.append({'carpeta': x['nombre'], 'otros': [y['nombre'] for y in hijos(nodos, x['ruta'], False)]})
    sueltos = [x['nombre'] for x in hijos(nodos, ruta, False) if not x['carpeta']]
    return salida, sueltos


def series_disco(nodos, ruta):
    salida = []
    for x in hijos(nodos, ruta):
        m = re.match(r'^(\d{3})\s+(.+)$', x['nombre'])
        temporadas = []
        for y in hijos(nodos, x['ruta']):
            mc = re.match(r'^(\d{1,2}\.\d)\s+(.+?)\s+(\d{1,2})ª', y['nombre'])
            if mc:
                temporadas.append({'codigo': mc.group(1), 'temporada': int(mc.group(3)), 'carpeta': y['nombre']})
        if m:
            salida.append({'numero': int(m.group(1)), 'serie': m.group(2).strip(), 'codigos': temporadas})
        else:  # Sello Vertigo / Sello Wildstorm
            salida.append({'sello': x['nombre'], 'series': [
                {'serie': y['nombre'], 'codigos': [
                    {'codigo': mc.group(1), 'temporada': int(mc.group(3))}
                    for z in hijos(nodos, y['ruta'])
                    if (mc := re.match(r'^(\d{1,2}\.\d)\s+(.+?)\s+(\d{1,2})ª', z['nombre']))]}
                for y in hijos(nodos, x['ruta']) if not y['nombre'].startswith('The Boys -') and not y['nombre'].startswith('The Boys 1x')]})
    return salida


def leer_animacion(carpeta):
    salida = []
    for sub in sorted(Path(carpeta).iterdir()):
        if not re.fullmatch(r'\d{3}', sub.name):
            continue
        for f in sub.glob('*.docx'):
            pars = [limpiar(p.text) for p in docx.Document(f).paragraphs if limpiar(p.text)]
            titulo = re.sub(r'\[editar\]$', '', pars[0]).strip() if pars else f.stem
            texto = ' '.join(pars[1:])
            fecha = re.search(r'(\d{1,2}) de (\w+) (?:del?|de) (\d{4})', texto)
            salida.append({'numero': int(sub.name), 'titulo': titulo, 'archivo': f.name,
                           'fecha_texto': fecha.group(0) if fecha else None,
                           'anio': int(fecha.group(3)) if fecha else None,
                           'resumen_fuente': texto})
    return salida


def main(originales, extraido):
    o, x = Path(originales), Path(extraido)
    img = x / 'imagenes'
    az = o / 'Archivador AZUL'
    cat = {}

    for clave, patron, pref in (('pre_dcue', 'PRE-DCUE Página *.docx', 'udc'), ('dcue', 'DCUE Página *.docx', 'dcdu')):
        filas = []
        for f in sorted(az.glob(patron)):
            if clave == 'dcue' and f.name.startswith('PRE'):
                continue
            _, fs = leer_pagina_indice(f, img / pref, pref)
            filas += fs
        cat[clave] = filas

    solo = {}
    for f in sorted(az.glob('word pelis solo *.docx')):
        heroe = re.match(r'word pelis solo (.+?) pag', f.name).group(1).strip().lower()
        _, fs = leer_pagina_indice(f, img / f'solo-{heroe}', f'solo-{heroe.replace(" ", "-")}')
        solo.setdefault(heroe, []).extend(fs)
    cat['pelis_solo'] = solo

    bloques = []
    for f in sorted(az.glob('series pagina *.docx'), key=lambda p: p.name):
        if '010' in f.name:  # copia de la 001 en preparación
            avisos.append('series pagina 010: es una copia de la página 001 a medio hacer; se ignora')
            continue
        bloques += leer_pagina_series(f, img / 'series')
    cat['series_paginas'] = bloques

    cat['resumen_carpetas'], cat['en_espera'] = leer_resumen(az / 'RESUMEN de carpetas UDC.docx')

    nodos = leer_estructura(o / 'ESTRUCTURA-DEL-DISCO.txt')
    base = 'TODO UDC/UDC Películas por orden'
    cat['disco_udc'], cat['disco_udc_sueltos'] = peliculas_disco(nodos, f'{base}/001 UDC')
    cat['disco_dcdu'], _ = peliculas_disco(nodos, f'{base}/002 DCDU')
    cat['disco_series'] = series_disco(nodos, 'UDC Series por orden')
    cat['disco_animacion'] = [x['nombre'] for x in hijos(nodos, 'animacion', False)] + \
        [x['nombre'] for x in hijos(nodos, 'animacion/Animación', False) if not re.fullmatch(r'\d{3}', x['nombre'])]

    cat['animacion'] = leer_animacion(o / 'animacion' / 'Animación')

    x.mkdir(parents=True, exist_ok=True)
    (x / 'catalogo.json').write_text(json.dumps(cat, ensure_ascii=False, indent=1), encoding='utf-8', newline='\n')
    (x / 'avisos.txt').write_text('\n'.join(avisos) + '\n', encoding='utf-8', newline='\n')
    print(f"PRE-DCUE {len(cat['pre_dcue'])} · DCUE {len(cat['dcue'])} · pelis solo {sum(map(len, solo.values()))} · "
          f"bloques de series {len(bloques)} · carpetas {len(cat['resumen_carpetas'])} · en espera {len(cat['en_espera'])} · "
          f"disco UDC {len(cat['disco_udc'])} / DCDU {len(cat['disco_dcdu'])} / series {len(cat['disco_series'])} · "
          f"animación {len(cat['animacion'])} · avisos {len(avisos)}")


if __name__ == '__main__':
    main(*sys.argv[1:3])
