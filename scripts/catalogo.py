"""
Reúne el catálogo completo de Marvel cruzando tus documentos, para incluir también
los títulos que aún no tienen ficha en los PowerPoint.

Uso:
    .venv/Scripts/python.exe scripts/catalogo.py <carpeta_originales_marvel> <carpeta_extraido_marvel>

Fuentes, por prioridad:
    1. 001 Índices.pptx         → UMC por orden cronológico y Pre-UMC por orden de estreno
    2. listado-videoteca.txt    → números que el índice aún no tiene (UMC 090-100, Pre-UMC 067-068)
    3. EL INDICE.docx           → series de imagen real Pre-UMC, películas y series de animación, LEGO

Genera `catalogo.json`: una entrada por título, con la fuente de cada dato.
"""
import json
import re
import sys
from pathlib import Path

from docx import Document
from pptx import Presentation

RE_ENTRADA = re.compile(r'^(\d{3})\s+(.*?)\((\d{4})\)\s*(.*)$', re.DOTALL)
TIPOS_INDICE = {'PELÍCULA': 'pelicula', 'SERIE': 'serie', 'ONE-SHOT': 'one-shot', 'SLINGSHOT': 'especial'}
# Orden importante: "cuatrilogía" contiene "trilogía".
GRUPOS = {'cuatrilogía': 'Cuatrilogía', 'trilogía': 'Trilogía', 'trilogia': 'Trilogía', 'saga': 'SAGA',
          'dulogía': 'Dulogía', 'corto': 'Corto', 'versión': 'Versión'}


def limpiar(t):
    t = (t or '').replace('\x0b', '\n').replace(' ', ' ').replace('\t', ' ')
    return '\n'.join(re.sub(r' {2,}', ' ', l).strip() for l in t.split('\n')).strip()


def una_linea(t):
    return ' '.join(l for l in limpiar(t).split('\n') if l).replace(' :', ':').strip()


def leer_indice(ruta):
    """Fichas del índice: tablas de 2 filas '001 (2011) / Título' + 'PELÍCULA'."""
    prs = Presentation(ruta)
    entradas, vistos = [], set()
    for n, diapo in enumerate(prs.slides, 1):
        textos = ' '.join(limpiar(s.text_frame.text) for s in diapo.shapes if s.has_text_frame)
        seccion = 'pre-umc' if 'PRE-UMC' in textos else ('umc' if 'cronológico' in textos else None)
        if not seccion:
            continue
        for s in diapo.shapes:
            if not getattr(s, 'has_table', False) or not s.has_table or len(s.table.rows) < 2:
                continue
            celdas = [limpiar(c.text) for c in s.table.rows[0].cells if limpiar(c.text)]
            tipo_txt = ' '.join(limpiar(c.text) for f in list(s.table.rows)[1:] for c in f.cells).strip().upper()
            if not celdas:
                continue
            m = RE_ENTRADA.match(celdas[0])
            if not m:
                continue
            numero, etiqueta, anio, resto = int(m.group(1)), m.group(2), int(m.group(3)), m.group(4)
            titulo = una_linea(resto.replace('/', '\n'))
            if len(titulo) <= 2 or (seccion, numero) in vistos:
                continue  # plantillas de relleno ("E", "I", "T"…) o repeticiones al final
            vistos.add((seccion, numero))
            e = {'seccion': seccion, 'numero': numero, 'anio': anio, 'titulo': titulo,
                 'tipo': TIPOS_INDICE.get(tipo_txt.replace(' ', ''), 'pelicula'),
                 'fuente': f'001 Índices.pptx · diapositiva {n}'}
            grupo = next((v for k, v in GRUPOS.items() if k in etiqueta.lower()), None)
            if grupo:
                e['grupoIndice'] = grupo
            # "Agentes SHIELD 1ª temp / epis 1-7" → temporada y tramo
            mt = re.search(r'(\d+)\s*[ªº]?\s*temp\w*', titulo, re.IGNORECASE)
            if mt:
                e['tipo'] = 'serie'  # el índice marca alguna temporada como PELÍCULA por error
                e['temporada'] = int(mt.group(1))
                e['titulo'] = re.split(r'\s*\d+\s*[ªº]?\s*temp', titulo, flags=re.IGNORECASE)[0].strip(' -/')
                me = re.search(r'epis\s*(\d+)\s*-\s*(\d+)', titulo, re.IGNORECASE)
                if me:
                    e['tramo'] = [int(me.group(1)), int(me.group(2))]
            entradas.append(e)
    return entradas


def leer_videoteca(ruta):
    """Carpetas '000 UMC\\091 PELI DEADPOOL y LOBEZNO' y '999 PRE UMC\\PELÍCULAS\\067 2024 Madame Web'."""
    umc, pre = {}, {}
    for linea in ruta.read_text(encoding='utf-8-sig').splitlines():
        partes = linea.strip().rstrip('\\').split('\\')
        if len(partes) == 4 and partes[2] == '000 UMC':
            m = re.match(r'(\d{3})\s+(.*)', partes[3])
            if m:
                umc[int(m.group(1))] = m.group(2)
        elif len(partes) == 5 and partes[2] == '999 PRE UMC' and partes[3] == 'PELÍCULAS':
            m = re.match(r'(\d{3})\s+(?:(\d{4})\s+)?(.*)', partes[4])
            if m:
                pre[int(m.group(1))] = (int(m.group(2)) if m.group(2) else None, m.group(3))
    return umc, pre


def entrada_videoteca_umc(numero, nombre):
    """'091 PELI DEADPOOL y LOBEZNO' → tipo y título; '094 serie Daredevil. Born Again (4) 1x completa'."""
    n = re.sub(r'^(COLOCAR|-+)\s*', '', nombre).strip()
    m = re.match(r'(PELI|serie|CORTO|one shot|slingshot)\s*(.*)', n, re.IGNORECASE)
    clase, resto = (m.group(1).lower(), m.group(2)) if m else ('', n)
    tipo = {'peli': 'pelicula', 'serie': 'serie', 'corto': 'corto', 'one shot': 'one-shot', 'slingshot': 'especial'}.get(clase, 'pelicula')
    e = {'seccion': 'umc', 'numero': numero, 'tipo': 'especial' if tipo == 'corto' else tipo,
         'titulo': resto.strip(), 'fuente': f'videoteca · {numero:03d} {nombre}'}
    mt = re.search(r'(\d+)x\s*(?:completa|comp)', resto, re.IGNORECASE)
    if tipo == 'serie' and mt:
        e['temporada'] = int(mt.group(1))
        e['titulo'] = re.split(r'\s*\(\d\)|\s*\d+x', resto)[0].strip(' .')
    e['titulo'] = re.sub(r'\s*\((haciendo|comprobar)\)', '', e['titulo'], flags=re.IGNORECASE).strip()
    return e


def leer_el_indice(ruta):
    """Columnas del Word 'EL INDICE': series Pre-UMC, películas animadas, series animadas y LEGO."""
    doc = Document(ruta)
    celdas = []
    for tabla in doc.tables:
        for fila in tabla.rows:
            previa = None
            for c in fila.cells:
                t = limpiar(c.text)
                if t and t != previa:
                    celdas.append(t)
                previa = t
    series_pre, pelis_anim, series_anim, lego = {}, {}, {}, {}
    for i, c in enumerate(celdas):
        lineas = [l for l in c.split('\n') if l]
        cabeza = lineas[0] if lineas else ''
        titulo = una_linea('\n'.join(lineas[1:])) if len(lineas) > 1 else ''
        # Serie de imagen real Pre-UMC: "1977/1979- 2 tempo" + título
        m = re.match(r'^(\d{4})(?:\s*[/–-]\s*(\d{4}))?\s*[-–]\s*(\d+)?\s*te\w*', cabeza)
        if not m:
            m = re.match(r'^(\d{4})/(\d{4})\s*-\s*$', cabeza)
        if m and titulo:
            series_pre.setdefault(titulo, {'anio': int(m.group(1)), 'anioFin': int(m.group(2)) if m.group(2) else None,
                                           'temporadas': int(m.group(3)) if m.lastindex and m.lastindex >= 3 and m.group(3) else None})
            continue
        # LEGO: "2013" + "Lego Marvel Super Heroes: …"
        if re.fullmatch(r'\d{4}', cabeza) and titulo.lower().startswith('lego'):
            lego.setdefault(titulo, {'anio': int(cabeza)})
            continue
        # Película animada: "2006" + "Ultimate Avengers" (o "1981" + "Drácula")
        if re.fullmatch(r'\d{4}', cabeza) and titulo and not re.match(r'\d', titulo):
            pelis_anim.setdefault(titulo, {'anio': int(cabeza)})
            continue
        # Serie animada: celdas consecutivas "1992-1997" | "X-Men" | "5"
        if re.fullmatch(r'\d{4}(-\d{4}|-)?', c) and i + 1 < len(celdas):
            nombre = una_linea(celdas[i + 1])
            temporadas = celdas[i + 2] if i + 2 < len(celdas) and re.fullmatch(r'\d+', celdas[i + 2]) else None
            if nombre and not re.match(r'[\d/]', nombre) and len(lineas) == 1:
                anios = re.findall(r'\d{4}', c)
                series_anim.setdefault(nombre, {'anio': int(anios[0]), 'anioFin': int(anios[1]) if len(anios) > 1 else None,
                                                'temporadas': int(temporadas) if temporadas else None})
    return series_pre, pelis_anim, series_anim, lego


def main(originales, extraido):
    originales, extraido = Path(originales), Path(extraido)
    indice = leer_indice(originales / 'UMC MARVEL' / 'Los Power Point juntos' / '001 Índices.pptx')
    umc_vid, pre_vid = leer_videoteca(extraido / 'listado-videoteca.txt')
    catalogo = list(indice)
    en_indice = {(e['seccion'], e['numero']) for e in indice}
    avisos = []

    for n, nombre in sorted(umc_vid.items()):
        if ('umc', n) not in en_indice:
            catalogo.append(entrada_videoteca_umc(n, nombre))
    for n, (anio, nombre) in sorted(pre_vid.items()):
        e_ind = next((e for e in indice if e['seccion'] == 'pre-umc' and e['numero'] == n), None)
        if e_ind is None:
            catalogo.append({'seccion': 'pre-umc', 'numero': n, 'anio': anio, 'titulo': nombre, 'tipo': 'pelicula',
                             'fuente': f'videoteca · {n:03d} {anio or ""} {nombre}'.replace('  ', ' ')})
        elif anio and e_ind['anio'] != anio or re.sub(r'\W', '', nombre.lower())[:6] not in re.sub(r'\W', '', e_ind['titulo'].lower()):
            avisos.append(f'Pre-UMC {n:03d}: índice "{e_ind["titulo"]}" ({e_ind["anio"]}) ≠ videoteca "{nombre}" ({anio})')

    series_pre, pelis_anim, series_anim, lego = leer_el_indice(originales / 'TODO WORD' / 'EL INDICE.docx')
    for titulo, d in series_pre.items():
        catalogo.append({'seccion': 'pre-umc', 'tipo': 'serie', 'titulo': titulo, **{k: v for k, v in d.items() if v},
                         'fuente': 'EL INDICE.docx · series Pre-UMC'})
    for titulo, d in pelis_anim.items():
        catalogo.append({'seccion': 'animacion', 'tipo': 'pelicula', 'animada': True, 'titulo': titulo, **d,
                         'fuente': 'EL INDICE.docx · películas de animación'})
    for titulo, d in series_anim.items():
        catalogo.append({'seccion': 'animacion', 'tipo': 'serie', 'animada': True, 'titulo': titulo,
                         **{k: v for k, v in d.items() if v}, 'fuente': 'EL INDICE.docx · series animadas'})
    for titulo, d in lego.items():
        catalogo.append({'seccion': 'animacion', 'tipo': 'especial', 'animada': True, 'titulo': titulo, **d,
                         'grupoIndice': 'LEGO', 'fuente': 'EL INDICE.docx · LEGO'})

    (extraido / 'catalogo.json').write_text(json.dumps(catalogo, ensure_ascii=False, indent=2), encoding='utf-8')
    (extraido / 'catalogo-avisos.txt').write_text('\n'.join(avisos) + '\n', encoding='utf-8')
    from collections import Counter
    print(len(catalogo), 'entradas:', dict(Counter((e['seccion'], e['tipo']) for e in catalogo)))
    print(len(avisos), 'diferencias índice/videoteca (ver catalogo-avisos.txt)')


if __name__ == '__main__':
    main(*sys.argv[1:3])
