"""
Extrae las fichas de los PowerPoint de Marvel a datos estructurados (JSON intermedio).

Uso:
    .venv/Scripts/python.exe scripts/extraer_fichas.py <carpeta_power_points> <carpeta_salida>

Lee (sin modificarlos):
    002 Fichas PELÍCULAS UMC.pptx       → películas del UMC
    003 Fichas SERIES UMC.pptx          → temporadas (presentación + páginas de capítulos)
    004 PRE-UMC Fichas PELÍCULAS.pptx   → películas Pre-UMC

Genera en la salida:
    fichas.json      → una entrada por ficha, con `avisos` cuando algo es dudoso
    imagenes/<hash>  → imágenes de las fichas (sin duplicados)
"""
import hashlib
import json
import re
import sys
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

DOCUMENTOS = [
    ('002 Fichas PELÍCULAS UMC.pptx', 'umc', 'pelicula'),
    ('003 Fichas SERIES UMC.pptx', 'umc', 'serie'),
    ('004 PRE-UMC Fichas PELÍCULAS.pptx', 'pre-umc', 'pelicula'),
]
# Una imagen que aparece en más fichas que esto es decorativa (logos, marcos).
UMBRAL_DECORATIVA = 4
# Insignias de la calidad de la copia (por el inicio de su huella sha1).
INSIGNIAS_CALIDAD = {'0b38fcd8a9': '720p', '97171bf5c3': '1080p', '971dc0bfe2': '4k', 'f59225962c': 'webm'}


def calidad_de(imagen):
    return next((c for h, c in INSIGNIAS_CALIDAD.items() if imagen['hash'].startswith(h)), None)


# ---------------------------------------------------------------- utilidades
def limpiar(texto):
    texto = (texto or '').replace('\x0b', '\n').replace(' ', ' ').replace('\t', ' ')
    lineas = [re.sub(r' {2,}', ' ', l).strip() for l in texto.split('\n')]
    return '\n'.join(lineas).strip()


def lineas(texto):
    return [l for l in limpiar(texto).split('\n') if l]


def una_linea(texto, sep=' / '):
    return sep.join(lineas(texto))


def cm(valor):
    return round(Emu(valor).cm, 1) if valor is not None else None


def minutos(texto):
    m = re.search(r'(\d+)\s*min', texto or '')
    return int(m.group(1)) if m else None


def numero(texto):
    m = re.fullmatch(r'\s*(\d{1,3})\s*', texto or '')
    return int(m.group(1)) if m else None


def sin_etiqueta(texto, *etiquetas):
    """Quita 'ARGUMENTO:' / 'Dirección' / 'Guión'… del principio de una celda."""
    t = limpiar(texto)
    for e in etiquetas:
        t = re.sub(rf'^{e}\s*:?\s*', '', t, flags=re.IGNORECASE)
    return t.strip()


def filas_tabla(forma):
    filas = []
    for fila in forma.table.rows:
        celdas = []
        for c in fila.cells:
            t = limpiar(c.text)
            if not celdas or celdas[-1] != t:  # las celdas combinadas se repiten
                celdas.append(t)
        filas.append(celdas)
    return filas


def formas_planas(formas):
    for f in formas:
        if f.shape_type == 6:  # grupo
            yield from formas_planas(f.shapes)
        else:
            yield f


# ---------------------------------------------------------------- lectura de diapositivas
def leer_diapositiva(diapo):
    """Devuelve tablas (con posición), textos e imágenes de una diapositiva."""
    tablas, textos, imagenes = [], [], []
    for f in formas_planas(diapo.shapes):
        pos = (cm(f.left), cm(f.top))
        if getattr(f, 'has_table', False) and f.has_table:
            tablas.append({'pos': pos, 'filas': filas_tabla(f)})
        elif f.shape_type == 13:
            blob = f.image.blob
            imagenes.append({
                'hash': hashlib.sha1(blob).hexdigest()[:16],
                'ext': f.image.ext,
                'blob': blob,
                'pos': pos,
                'tam': (cm(f.width), cm(f.height)),
            })
        elif f.has_text_frame and f.text_frame.text.strip():
            textos.append({'pos': pos, 'texto': limpiar(f.text_frame.text)})
    notas = ''
    if diapo.has_notes_slide:
        notas = limpiar(diapo.notes_slide.notes_text_frame.text)
    return {'tablas': tablas, 'textos': textos, 'imagenes': imagenes, 'notas': notas}


def es_ficha(d, clase):
    return any(f'Ficha Técnica\n{clase}' in t['texto'] or f'Ficha Técnica {clase}' in t['texto'] for t in d['textos'])


def tabla_que_empieza(d, primera):
    for t in d['tablas']:
        if t['filas'] and t['filas'][0] and t['filas'][0][0].lower().startswith(primera.lower()):
            return t
    return None


def pares_clave_valor(tabla):
    """Tabla de 2 columnas → dict. Reconoce también la fila '001 | Orden cronológico'."""
    datos = {}
    for fila in tabla['filas']:
        if len(fila) < 2:
            continue
        a, b = fila[0], fila[1]
        if numero(a) is not None and 'orden' in b.lower():
            datos['orden'] = a
        else:
            datos[a.lower().replace('ó', 'o').replace('ñ', 'n').strip()] = b
    return datos


def leer_reparto(d):
    t = tabla_que_empieza(d, 'Actor')
    if not t:
        return []
    reparto = []
    for fila in t['filas'][1:]:
        if fila and fila[0]:
            reparto.append({'actor': una_linea(fila[0], ' '), 'personaje': una_linea(fila[1]) if len(fila) > 1 else ''})
    return reparto


def comunes(d, kv, avisos):
    """Campos que comparten películas y temporadas."""
    r = {}
    if 'orden' in kv or 'orden cronologico' in kv:
        r['numero'] = numero(kv.get('orden') or kv.get('orden cronologico'))
    pagina = re.search(r'\d+', kv.get('indice', ''))
    if pagina:
        r['paginaArchivador'] = int(pagina.group())
    anio = re.search(r'(19|20)\d{2}', kv.get('ano', ''))
    if anio:
        r['anio'] = int(anio.group())
    else:
        avisos.append('Sin año')
    for clave, campo in [('nacionalidad', 'nacionalidad')]:
        if kv.get(clave):
            r[campo] = una_linea(kv[clave])
    for clave, campo in [('direccion', 'direccion'), ('guion', 'guion'), ('productora', 'productora'), ('distribuidora', 'distribuidora')]:
        if kv.get(clave):
            r[campo] = lineas(kv[clave])
    if kv.get('duracion'):
        r['duracion'] = minutos(kv['duracion'])
    r['reparto'] = leer_reparto(d)
    arg = tabla_que_empieza(d, 'ARGUMENTO')
    if arg:
        r['argumento'] = sin_etiqueta(arg['filas'][0][0], 'ARGUMENTO')
    else:
        avisos.append('Sin argumento')
    return r


# ---------------------------------------------------------------- películas
def ficha_pelicula(d, seccion):
    avisos = []
    kv_tabla = tabla_que_empieza(d, 'Índice')
    if not kv_tabla:
        return None
    kv = pares_clave_valor(kv_tabla)
    r = {'tipo': 'pelicula', 'seccion': seccion}
    # El título es la tabla de una sola celda situada encima de los datos técnicos.
    candidatos = [t for t in d['tablas'] if len(t['filas']) == 1 and len(t['filas'][0]) == 1 and t['filas'][0][0]
                  and not t['filas'][0][0].upper().startswith('ARGUMENTO')]
    if candidatos:
        r['titulo'] = una_linea(candidatos[0]['filas'][0][0], ' ')
    else:
        avisos.append('Sin título')
    r.update(comunes(d, kv, avisos))
    r['avisos'] = avisos
    return r


# ---------------------------------------------------------------- series
RE_TEMPORADA = re.compile(r'(\d+)\s*[ªº°]?\s*temporada', re.IGNORECASE)
RE_TRAMO = re.compile(r'del\s+(\d+)\s+al\s+(\d+)', re.IGNORECASE)
RE_CAPITULO = re.compile(r'^(\d+)x(\d{2,3})\s*-?\s*(.*)$', re.DOTALL)


def presentacion_temporada(d):
    """Diapositiva 'Ficha Técnica SERIE' con la tabla Índice / Orden Cronológico."""
    kv_tabla = tabla_que_empieza(d, 'Índice')
    if not kv_tabla:
        return None
    avisos = []
    kv = pares_clave_valor(kv_tabla)
    r = {'tipo': 'serie', 'seccion': 'umc'}
    cabecera = next((t for t in d['tablas'] if len(t['filas']) == 1 and len(t['filas'][0]) == 2
                     and RE_TEMPORADA.search(t['filas'][0][1])), None)
    if cabecera:
        r['titulo'] = una_linea(cabecera['filas'][0][0], ' ')
        m = RE_TEMPORADA.search(cabecera['filas'][0][1])
        r['temporada'] = int(m.group(1))
        tramo = RE_TRAMO.search(cabecera['filas'][0][1])
        if tramo:
            r['tramo'] = (int(tramo.group(1)), int(tramo.group(2)))
    else:
        avisos.append('Sin título/temporada en la presentación')
    r.update(comunes(d, kv, avisos))
    lista = tabla_que_empieza(d, 'Capítulo')
    r['listaCapitulos'] = []
    if lista:
        for fila in lista['filas'][1:]:
            if len(fila) >= 2 and re.fullmatch(r'\d+x\d+', fila[0]):
                r['listaCapitulos'].append({'codigo': fila[0], 'titulo': una_linea(fila[1], ' ').rstrip('.')})
    r['capitulos'] = []
    r['avisos'] = avisos
    return r


def capitulos_de_pagina(d):
    """Tablas de capítulo: fila 1 '1x01- Título', fila 2 'ARGUMENTO …', fila 3 'Dirección | Guión | Duración'."""
    caps = []
    for t in sorted(d['tablas'], key=lambda t: (t['pos'][1] or 0, t['pos'][0] or 0)):
        filas = t['filas']
        if len(filas) < 2 or not filas[0]:
            continue
        m = RE_CAPITULO.match(filas[0][0])
        if not m:
            continue
        cap = {'codigo': f'{int(m.group(1))}x{m.group(2)}', 'titulo': una_linea(m.group(3), ' ').strip(' .-')}
        for fila in filas[1:]:
            for celda in fila:
                bajo = celda.lower()
                if bajo.startswith('argumento'):
                    cap['argumento'] = sin_etiqueta(celda, 'ARGUMENTO')
                elif bajo.startswith('direcci'):
                    cap['direccion'] = lineas(sin_etiqueta(celda, 'Dirección', 'Direccion'))
                elif bajo.startswith('gui'):
                    cap['guion'] = lineas(sin_etiqueta(celda, 'Guión', 'Guion'))
                elif bajo.startswith('duraci'):
                    cap['duracion'] = minutos(celda)
        caps.append(cap)
    return caps


# ---------------------------------------------------------------- principal
def main(carpeta, salida):
    carpeta, salida = Path(carpeta), Path(salida)
    (salida / 'imagenes').mkdir(parents=True, exist_ok=True)

    fichas, borradores, actual = [], [], None
    for nombre, seccion, clase in DOCUMENTOS:
        prs = Presentation(carpeta / nombre)
        for i, diapo in enumerate(prs.slides, 1):
            d = leer_diapositiva(diapo)
            fuente = f'{nombre} · diapositiva {i}'
            if clase == 'pelicula':
                if not es_ficha(d, 'PELÍCULA'):
                    continue
                f = ficha_pelicula(d, seccion)
                if f:
                    f.update(fuente=fuente, notasDiapositiva=d['notas'], _imagenes=list(d['imagenes']))
                    fichas.append(f)
                continue
            # Series: una presentación abre temporada; las páginas siguientes aportan capítulos.
            if not es_ficha(d, 'SERIE'):
                continue
            pres = presentacion_temporada(d)
            if pres:
                clave = (pres.get('titulo'), pres.get('temporada'), pres.get('numero'), pres.get('tramo'))
                if actual is not None and clave == actual['_clave'] and not actual['capitulos']:
                    # La lista de capítulos continúa en una segunda diapositiva de presentación.
                    actual['listaCapitulos'].extend(pres['listaCapitulos'])
                    actual['fuentesExtra'].append(fuente)
                    actual['_imagenes'].extend(d['imagenes'])
                    continue
                pres.update(fuente=fuente, notasDiapositiva=d['notas'], _imagenes=list(d['imagenes']),
                            _clave=clave, fuentesExtra=[])
                fichas.append(pres)
                actual = pres
                continue
            caps = capitulos_de_pagina(d)
            if caps and actual is not None:
                actual['capitulos'].extend(caps)
                actual['fuentesExtra'].append(fuente)
                actual['_imagenes'].extend(d['imagenes'])
            elif caps:
                print(f'Aviso: capítulos sin temporada en {fuente}')

    # Borradores: plantillas a medio rellenar ("A: Q.", título vacío…). Se apartan, no se borran.
    def es_borrador(f):
        t = f.get('titulo', '')
        return not t or re.fullmatch(r'[A-ZÁÉÍÓÚ]: ?[A-ZÁÉÍÓÚ]\.?', t) or (f['tipo'] == 'serie' and 'temporada' not in f)

    for f in list(fichas):
        if es_borrador(f):
            fichas.remove(f)
            borradores.append({'fuente': f['fuente'], 'titulo': f.get('titulo'), 'numero': f.get('numero'),
                               'notasDiapositiva': f['notasDiapositiva']})

    # Imágenes decorativas: las que se repiten en muchas fichas distintas (logos, marcos).
    usos = Counter(h for f in fichas for h in {im['hash'] for im in f['_imagenes']})
    for f in fichas:
        propias, vistas, calidad = [], set(), []
        for im in f.pop('_imagenes'):
            c = calidad_de(im)
            if c:
                if c not in calidad:
                    calidad.append(c)
                continue
            if usos[im['hash']] > UMBRAL_DECORATIVA or im['hash'] in vistas:
                continue
            vistas.add(im['hash'])
            archivo = f"{im['hash']}.{im['ext']}"
            destino = salida / 'imagenes' / archivo
            if not destino.exists():
                destino.write_bytes(im['blob'])
            propias.append({'archivo': archivo, 'pos': im['pos'], 'tam': im['tam']})
        f['imagenes'] = propias
        f['calidad'] = calidad
        f.pop('_clave', None)

    for f in fichas:
        if f['tipo'] != 'serie':
            continue
        # Lista de capítulos con la temporada equivocada (p. ej. 1x01 en la 2ª temporada).
        mal = [c for c in f['listaCapitulos'] if int(c['codigo'].split('x')[0]) != f['temporada']]
        if mal and len(mal) == len(f['listaCapitulos']):
            original = mal[0]['codigo'].split('x')[0]
            for c in f['listaCapitulos']:
                c['codigo'] = f"{f['temporada']}x{c['codigo'].split('x')[1]}"
            f['avisos'].append(f'Corregido: la lista de capítulos usaba códigos {original}x en la temporada {f["temporada"]}')
        vistos = {c['codigo'] for c in f['capitulos']}
        faltan = [c['codigo'] for c in f['listaCapitulos'] if c['codigo'] not in vistos]
        if faltan:
            f['avisos'].append(f'Capítulos en la lista sin página de detalle: {", ".join(faltan)}')
        sobran = [c['codigo'] for c in f['capitulos'] if c['codigo'] not in {x['codigo'] for x in f['listaCapitulos']}]
        if sobran and f['listaCapitulos']:
            f['avisos'].append(f'Capítulos con página pero no en la lista: {", ".join(sobran)}')

    (salida / 'fichas.json').write_text(json.dumps(fichas, ensure_ascii=False, indent=2), encoding='utf-8')
    (salida / 'borradores.json').write_text(json.dumps(borradores, ensure_ascii=False, indent=2), encoding='utf-8')
    por_tipo = Counter((f['seccion'], f['tipo']) for f in fichas)
    print(f'{len(fichas)} fichas extraídas:', dict(por_tipo))
    print(f'{len(borradores)} borradores apartados')
    print(f"Imágenes propias guardadas: {len(list((salida / 'imagenes').iterdir()))}")
    print(f"Fichas con avisos: {sum(1 for f in fichas if f['avisos'])}")


if __name__ == '__main__':
    main(*sys.argv[1:3])
